import os
import io
from datetime import datetime
from typing import Dict, Any
from flask import current_app
try:
    from weasyprint import HTML, CSS
    WEASYPRINT_AVAILABLE = True
except ImportError:
    WEASYPRINT_AVAILABLE = False

try:
    from docx import Document
    from docx.shared import Inches
    PYTHON_DOCX_AVAILABLE = True
except ImportError:
    PYTHON_DOCX_AVAILABLE = False

class PDFService:
    """
    PDF generation service for legal documents.
    Supports both PDF and Word document generation.
    """

    def __init__(self):
        self.template_dir = os.path.join(current_app.root_path, 'templates')
        self.output_dir = current_app.config.get('UPLOAD_FOLDER', 'uploads')

    def generate_fir_pdf(self, fir_data: Dict[str, Any], output_filename: str = None) -> str:
        """
        Generate PDF for FIR document.

        Args:
            fir_data (Dict): FIR data including user info, content, etc.
            output_filename (str): Optional output filename

        Returns:
            str: Path to generated PDF file
        """
        if not WEASYPRINT_AVAILABLE:
            return self._generate_text_fir(fir_data, output_filename)

        try:
            # Generate HTML content
            html_content = self._generate_fir_html(fir_data)

            # Generate CSS
            css_content = self._get_fir_css()

            # Create HTML object
            html = HTML(string=html_content)
            css = CSS(string=css_content)

            # Generate output filename
            if not output_filename:
                timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
                output_filename = f"FIR_{timestamp}.pdf"

            # Ensure output directory exists
            os.makedirs(self.output_dir, exist_ok=True)

            # Generate PDF
            output_path = os.path.join(self.output_dir, output_filename)
            html.write_pdf(output_path, stylesheets=[css])

            return output_path

        except Exception as e:
            current_app.logger.error(f"PDF generation failed: {str(e)}")
            # Fallback to text format
            return self._generate_text_fir(fir_data, output_filename)

    def generate_fir_word(self, fir_data: Dict[str, Any], output_filename: str = None) -> str:
        """
        Generate Word document for FIR.

        Args:
            fir_data (Dict): FIR data
            output_filename (str): Optional output filename

        Returns:
            str: Path to generated Word document
        """
        if not PYTHON_DOCX_AVAILABLE:
            raise Exception("python-docx not available for Word generation")

        try:
            # Create new document
            doc = Document()

            # Add title
            title = doc.add_heading('FIRST INFORMATION REPORT', 0)
            title.alignment = 1  # Center

            # Add FIR details
            self._add_fir_content_to_doc(doc, fir_data)

            # Generate output filename
            if not output_filename:
                timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
                output_filename = f"FIR_{timestamp}.docx"

            # Ensure output directory exists
            os.makedirs(self.output_dir, exist_ok=True)

            # Save document
            output_path = os.path.join(self.output_dir, output_filename)
            doc.save(output_path)

            return output_path

        except Exception as e:
            current_app.logger.error(f"Word document generation failed: {str(e)}")
            raise e

    def generate_complaint_pdf(self, complaint_data: Dict[str, Any], output_filename: str = None) -> str:
        """
        Generate PDF for complaint document.

        Args:
            complaint_data (Dict): Complaint data
            output_filename (str): Optional output filename

        Returns:
            str: Path to generated PDF file
        """
        if not WEASYPRINT_AVAILABLE:
            return self._generate_text_complaint(complaint_data, output_filename)

        try:
            # Generate HTML content
            html_content = self._generate_complaint_html(complaint_data)

            # Generate CSS
            css_content = self._get_complaint_css()

            # Create HTML object
            html = HTML(string=html_content)
            css = CSS(string=css_content)

            # Generate output filename
            if not output_filename:
                timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
                output_filename = f"Complaint_{timestamp}.pdf"

            # Ensure output directory exists
            os.makedirs(self.output_dir, exist_ok=True)

            # Generate PDF
            output_path = os.path.join(self.output_dir, output_filename)
            html.write_pdf(output_path, stylesheets=[css])

            return output_path

        except Exception as e:
            current_app.logger.error(f"Complaint PDF generation failed: {str(e)}")
            return self._generate_text_complaint(complaint_data, output_filename)

    def _generate_fir_html(self, fir_data: Dict[str, Any]) -> str:
        """Generate HTML content for FIR."""
        user_info = fir_data.get('user_info', {})
        fir_content = fir_data.get('fir_content', {})
        suggested_sections = fir_data.get('suggested_sections', [])

        html = f"""
        <!DOCTYPE html>
        <html>
        <head>
            <meta charset="utf-8">
            <title>First Information Report</title>
        </head>
        <body>
            <div class="document">
                <header class="header">
                    <h1>FIRST INFORMATION REPORT</h1>
                    <div class="fir-number">FIR No: {fir_content.get('id', 'N/A')}</div>
                </header>

                <div class="content">
                    <section class="section">
                        <h2>Police Information</h2>
                        <div class="info-grid">
                            <div class="info-item">
                                <span class="label">Police Station:</span>
                                <span class="value">{fir_content.get('police_station', 'To be filled')}</span>
                            </div>
                            <div class="info-item">
                                <span class="label">District:</span>
                                <span class="value">{fir_content.get('district', 'To be filled')}</span>
                            </div>
                            <div class="info-item">
                                <span class="label">Date of Report:</span>
                                <span class="value">{datetime.now().strftime('%d/%m/%Y')}</span>
                            </div>
                            <div class="info-item">
                                <span class="label">Time of Report:</span>
                                <span class="value">{datetime.now().strftime('%H:%M:%S')}</span>
                            </div>
                        </div>
                    </section>

                    <section class="section">
                        <h2>Complainant Information</h2>
                        <div class="info-grid">
                            <div class="info-item">
                                <span class="label">Name:</span>
                                <span class="value">{user_info.get('name', 'N/A')}</span>
                            </div>
                            <div class="info-item">
                                <span class="label">Email:</span>
                                <span class="value">{user_info.get('email', 'N/A')}</span>
                            </div>
                            <div class="info-item">
                                <span class="label">Phone:</span>
                                <span class="value">{user_info.get('phone', 'N/A')}</span>
                            </div>
                            <div class="info-item">
                                <span class="label">Address:</span>
                                <span class="value">{user_info.get('address', 'N/A')}</span>
                            </div>
                        </div>
                    </section>

                    <section class="section">
                        <h2>Incident Details</h2>
                        <div class="info-grid">
                            <div class="info-item">
                                <span class="label">Date of Incident:</span>
                                <span class="value">{fir_content.get('incident_date', 'N/A')}</span>
                            </div>
                            <div class="info-item">
                                <span class="label">Time of Incident:</span>
                                <span class="value">{fir_content.get('incident_time', 'N/A')}</span>
                            </div>
                            <div class="info-item">
                                <span class="label">Location:</span>
                                <span class="value">{fir_content.get('incident_location', 'N/A')}</span>
                            </div>
                            <div class="info-item">
                                <span class="label">FIR Type:</span>
                                <span class="value">{fir_content.get('fir_type', 'N/A')}</span>
                            </div>
                        </div>
                    </section>

                    <section class="section">
                        <h2>Complaint Details</h2>
                        <div class="complaint-text">
                            {fir_content.get('complaint_text', 'No complaint text provided')}
                        </div>
                    </section>

                    {self._generate_sections_html(suggested_sections)}
                </div>

                <footer class="footer">
                    <div class="signature-section">
                        <div class="signature-box">
                            <p>Signature of Complainant</p>
                            <div class="signature-line"></div>
                        </div>
                        <div class="signature-box">
                            <p>Signature of Police Officer</p>
                            <div class="signature-line"></div>
                            <p>Rank: ___________________</p>
                            <p>Name: ___________________</p>
                        </div>
                    </div>
                    <div class="disclaimer">
                        <p><strong>Disclaimer:</strong> This FIR has been generated using AI assistance.
                        All information should be verified and cross-checked by the investigating officer.</p>
                    </div>
                </footer>
            </div>
        </body>
        </html>
        """
        return html

    def _generate_sections_html(self, sections: list) -> str:
        """Generate HTML for suggested legal sections."""
        if not sections:
            return ""

        html = """
        <section class="section">
            <h2>Suggested Legal Sections</h2>
            <div class="sections-list">
        """

        for section in sections[:5]:  # Limit to 5 sections
            html += f"""
                <div class="section-item">
                    <h3>{section.get('section', 'N/A')} - {section.get('title', 'N/A')}</h3>
                    <p><strong>Description:</strong> {section.get('description', 'N/A')}</p>
                    <p><strong>Punishment:</strong> {section.get('punishment', 'N/A')}</p>
                    <p><strong>Bailable:</strong> {section.get('bailable', 'N/A')}</p>
                    <p><strong>Cognizable:</strong> {section.get('cognizable', 'N/A')}</p>
                </div>
            """

        html += """
            </div>
        </section>
        """

        return html

    def _generate_complaint_html(self, complaint_data: Dict[str, Any]) -> str:
        """Generate HTML content for complaint."""
        user_info = complaint_data.get('user_info', {})
        complaint_content = complaint_data.get('complaint_content', {})
        formatted_content = complaint_data.get('formatted_content', '')

        html = f"""
        <!DOCTYPE html>
        <html>
        <head>
            <meta charset="utf-8">
            <title>Formal Complaint</title>
        </head>
        <body>
            <div class="document">
                <header class="header">
                    <h1>FORMAL COMPLAINT</h1>
                </header>

                <div class="content">
                    <section class="section">
                        <h2>Complainant Information</h2>
                        <div class="info-grid">
                            <div class="info-item">
                                <span class="label">Name:</span>
                                <span class="value">{user_info.get('name', 'N/A')}</span>
                            </div>
                            <div class="info-item">
                                <span class="label">Email:</span>
                                <span class="value">{user_info.get('email', 'N/A')}</span>
                            </div>
                            <div class="info-item">
                                <span class="label">Phone:</span>
                                <span class="value">{user_info.get('phone', 'N/A')}</span>
                            </div>
                            <div class="info-item">
                                <span class="label">Date:</span>
                                <span class="value">{datetime.now().strftime('%d/%m/%Y')}</span>
                            </div>
                        </div>
                    </section>

                    <section class="section">
                        <h2>Complaint Details</h2>
                        <div class="complaint-content">
                            {formatted_content or complaint_content.get('original_text', 'No complaint text provided')}
                        </div>
                    </section>
                </div>

                <footer class="footer">
                    <div class="signature-section">
                        <div class="signature-box">
                            <p>Signature of Complainant</p>
                            <div class="signature-line"></div>
                            <p>Date: ___________________</p>
                        </div>
                    </div>
                </footer>
            </div>
        </body>
        </html>
        """
        return html

    def _get_fir_css(self) -> str:
        """Get CSS styles for FIR document."""
        return """
        <style>
            @page {
                size: A4;
                margin: 2cm;
            }

            body {
                font-family: 'Times New Roman', serif;
                font-size: 12pt;
                line-height: 1.5;
                color: #000;
            }

            .document {
                max-width: 100%;
            }

            .header {
                text-align: center;
                margin-bottom: 30px;
                border-bottom: 2px solid #000;
                padding-bottom: 20px;
            }

            .header h1 {
                font-size: 18pt;
                font-weight: bold;
                margin: 0;
                text-transform: uppercase;
            }

            .fir-number {
                font-size: 14pt;
                font-weight: bold;
                margin-top: 10px;
            }

            .section {
                margin-bottom: 25px;
            }

            .section h2 {
                font-size: 14pt;
                font-weight: bold;
                margin-bottom: 15px;
                border-bottom: 1px solid #ccc;
                padding-bottom: 5px;
            }

            .info-grid {
                display: grid;
                grid-template-columns: 1fr 1fr;
                gap: 10px;
                margin-bottom: 20px;
            }

            .info-item {
                display: flex;
                margin-bottom: 8px;
            }

            .info-item .label {
                font-weight: bold;
                min-width: 120px;
                margin-right: 10px;
            }

            .info-item .value {
                flex: 1;
            }

            .complaint-text {
                background-color: #f9f9f9;
                padding: 15px;
                border: 1px solid #ddd;
                border-radius: 5px;
                white-space: pre-wrap;
                min-height: 100px;
            }

            .sections-list {
                margin-top: 15px;
            }

            .section-item {
                background-color: #f9f9f9;
                padding: 15px;
                border: 1px solid #ddd;
                border-radius: 5px;
                margin-bottom: 15px;
            }

            .section-item h3 {
                margin-top: 0;
                color: #1a365d;
            }

            .footer {
                margin-top: 40px;
                border-top: 1px solid #ccc;
                padding-top: 20px;
            }

            .signature-section {
                display: flex;
                justify-content: space-around;
                margin-bottom: 30px;
            }

            .signature-box {
                text-align: center;
                min-width: 200px;
            }

            .signature-line {
                border-bottom: 1px solid #000;
                height: 40px;
                margin: 10px 0;
            }

            .disclaimer {
                font-style: italic;
                font-size: 10pt;
                color: #666;
                text-align: center;
            }
        </style>
        """

    def _get_complaint_css(self) -> str:
        """Get CSS styles for complaint document."""
        return self._get_fir_css()  # Use same styling

    def _add_fir_content_to_doc(self, doc, fir_data: Dict[str, Any]):
        """Add FIR content to Word document."""
        user_info = fir_data.get('user_info', {})
        fir_content = fir_data.get('fir_content', {})

        # Add police information
        doc.add_heading('Police Information', level=1)
        p = doc.add_paragraph()
        p.add_run('Police Station: ').bold = True
        p.add_run(fir_content.get('police_station', 'To be filled'))

        p = doc.add_paragraph()
        p.add_run('District: ').bold = True
        p.add_run(fir_content.get('district', 'To be filled'))

        p = doc.add_paragraph()
        p.add_run('Date of Report: ').bold = True
        p.add_run(datetime.now().strftime('%d/%m/%Y'))

        # Add complainant information
        doc.add_heading('Complainant Information', level=1)
        p = doc.add_paragraph()
        p.add_run('Name: ').bold = True
        p.add_run(user_info.get('name', 'N/A'))

        p = doc.add_paragraph()
        p.add_run('Email: ').bold = True
        p.add_run(user_info.get('email', 'N/A'))

        # Add incident details
        doc.add_heading('Incident Details', level=1)
        p = doc.add_paragraph()
        p.add_run('Date of Incident: ').bold = True
        p.add_run(str(fir_content.get('incident_date', 'N/A')))

        p = doc.add_paragraph()
        p.add_run('Location: ').bold = True
        p.add_run(fir_content.get('incident_location', 'N/A'))

        # Add complaint text
        doc.add_heading('Complaint Details', level=1)
        doc.add_paragraph(fir_content.get('complaint_text', 'No complaint text provided'))

    def _generate_text_fir(self, fir_data: Dict[str, Any], output_filename: str = None) -> str:
        """Generate plain text FIR as fallback."""
        user_info = fir_data.get('user_info', {})
        fir_content = fir_data.get('fir_content', {})

        text_content = f"""
FIRST INFORMATION REPORT
========================

FIR No: {fir_content.get('id', 'N/A')}
Date: {datetime.now().strftime('%d/%m/%Y')}
Time: {datetime.now().strftime('%H:%M:%S')}

POLICE INFORMATION
------------------
Police Station: {fir_content.get('police_station', 'To be filled')}
District: {fir_content.get('district', 'To be filled')}

COMPLAINANT INFORMATION
-----------------------
Name: {user_info.get('name', 'N/A')}
Email: {user_info.get('email', 'N/A')}
Phone: {user_info.get('phone', 'N/A')}

INCIDENT DETAILS
----------------
Date of Incident: {fir_content.get('incident_date', 'N/A')}
Location: {fir_content.get('incident_location', 'N/A')}
FIR Type: {fir_content.get('fir_type', 'N/A')}

COMPLAINT DETAILS
-----------------
{fir_content.get('complaint_text', 'No complaint text provided')}

_________________________
Signature of Complainant

_________________________
Signature of Police Officer
Rank: ___________________
Name: ___________________
        """

        # Generate output filename
        if not output_filename:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            output_filename = f"FIR_{timestamp}.txt"

        # Ensure output directory exists
        os.makedirs(self.output_dir, exist_ok=True)

        # Save text file
        output_path = os.path.join(self.output_dir, output_filename)
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(text_content)

        return output_path

    def _generate_text_complaint(self, complaint_data: Dict[str, Any], output_filename: str = None) -> str:
        """Generate plain text complaint as fallback."""
        user_info = complaint_data.get('user_info', {})
        complaint_content = complaint_data.get('complaint_content', {})

        text_content = f"""
FORMAL COMPLAINT
===============

Date: {datetime.now().strftime('%d/%m/%Y')}

COMPLAINANT INFORMATION
-----------------------
Name: {user_info.get('name', 'N/A')}
Email: {user_info.get('email', 'N/A')}
Phone: {user_info.get('phone', 'N/A')}

COMPLAINT DETAILS
-----------------
{complaint_content.get('original_text', 'No complaint text provided')}

_________________________
Signature of Complainant
Date: ___________________
        """

        # Generate output filename
        if not output_filename:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            output_filename = f"Complaint_{timestamp}.txt"

        # Ensure output directory exists
        os.makedirs(self.output_dir, exist_ok=True)

        # Save text file
        output_path = os.path.join(self.output_dir, output_filename)
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(text_content)

        return output_path


# Global PDF service instance
pdf_service = PDFService()